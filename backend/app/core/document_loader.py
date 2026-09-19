import io
from pathlib import Path
from typing import Any, Dict, List, Union
import pymupdf as fitz  # PyMuPDF
import docx

from app.core.logging import logger
from app.core.normalizer import normalize_text


class DocumentLoader:
    """
    Unified extraction interface preserving page numbers and structural metadata.
    Protects against malformed payloads, non-PDF streams, and exhausted file buffers.
    """

    @classmethod
    def _to_bytes(cls, file_input: Union[str, Path, bytes, Any]) -> bytes:
        """
        Safely extracts raw bytes from paths, buffers, or Streamlit UploadedFile objects.
        Resets file pointers when seekable.
        """
        if hasattr(file_input, "seek") and hasattr(file_input, "read"):
            file_input.seek(0)
            data = file_input.read()
            file_input.seek(0)
        elif isinstance(file_input, (str, Path)):
            path = Path(file_input)
            if not path.exists():
                raise FileNotFoundError(f"File not found: {path.resolve()}")
            data = path.read_bytes()
        elif isinstance(file_input, (bytes, bytearray)):
            data = bytes(file_input)
        else:
            raise TypeError(f"Unsupported input type: {type(file_input)}")

        if not data or len(data) == 0:
            raise ValueError("Input file buffer contains 0 bytes.")

        return data

    @staticmethod
    def load_txt(file_input: Union[str, Path, bytes, Any]) -> List[Dict[str, Any]]:
        """Extracts text from a UTF-8 encoded text stream or file."""
        if isinstance(file_input, (str, Path)):
            with open(file_input, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif isinstance(file_input, (bytes, bytearray)):
            content = file_input.decode("utf-8", errors="replace")
        elif hasattr(file_input, "seek") and hasattr(file_input, "read"):
            file_input.seek(0)
            raw = file_input.read()
            file_input.seek(0)
            content = raw.decode("utf-8", errors="replace") if isinstance(raw, (bytes, bytearray)) else str(raw)
        else:
            raise TypeError(f"Unsupported TXT input type: {type(file_input)}")

        normalized = normalize_text(content)
        return [{"page": 1, "text": normalized, "blocks": []}]

    @staticmethod
    def load_docx(file_input: Union[str, Path, bytes, Any]) -> List[Dict[str, Any]]:
        """Extracts paragraphs and tables from a DOCX file."""
        if isinstance(file_input, (str, Path)):
            doc = docx.Document(str(file_input))
        elif isinstance(file_input, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_input))
        elif hasattr(file_input, "seek") and hasattr(file_input, "read"):
            file_input.seek(0)
            doc = docx.Document(file_input)
            file_input.seek(0)
        else:
            raise TypeError(f"Unsupported DOCX input type: {type(file_input)}")

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = normalize_text("\n\n".join(paragraphs))
        return [{"page": 1, "text": full_text, "blocks": []}]

    @classmethod
    def load_pdf(cls, file_input: Union[str, Path, bytes, Any]) -> List[Dict[str, Any]]:
        """
        Extracts pages from PDF using PyMuPDF.
        Validates magic header before sending stream to PyMuPDF C-parser.
        """
        if isinstance(file_input, (str, Path)):
            path = Path(file_input)
            if not path.exists() or path.stat().st_size == 0:
                raise ValueError(f"PDF file does not exist or is 0 bytes: {path.resolve()}")
            stream_bytes = path.read_bytes()
        else:
            stream_bytes = cls._to_bytes(file_input)

        if not stream_bytes.startswith(b"%PDF-"):
            header_peek = stream_bytes[:30]
            raise ValueError(
                f"Invalid PDF header format. Expected '%PDF-', but received {header_peek!r}. "
                "Ensure that non-PDF text files are not being routed through load_pdf."
            )

        doc = fitz.open(stream=stream_bytes, filetype="pdf")
        pages_data = []
        try:
            for page_index in range(len(doc)):
                page = doc[page_index]
                page_num = page_index + 1

                raw_text = page.get_text("text")
                normalized_text = normalize_text(raw_text)

                raw_blocks = page.get_text("blocks")
                blocks = []
                for b in raw_blocks:
                    if len(b) >= 5 and isinstance(b[4], str) and b[4].strip():
                        blocks.append({
                            "bbox": (round(b[0], 2), round(b[1], 2), round(b[2], 2), round(b[3], 2)),
                            "text": normalize_text(b[4]),
                        })

                pages_data.append({
                    "page": page_num,
                    "text": normalized_text,
                    "blocks": blocks,
                })
        finally:
            doc.close()

        return pages_data

    @classmethod
    def load_document(
        cls, file_input: Union[str, Path, bytes, Any], filename: str
    ) -> List[Dict[str, Any]]:
        """Routes file by extension."""
        suffix = Path(filename).suffix.lower()

        if suffix == ".pdf":
            return cls.load_pdf(file_input)
        elif suffix == ".docx":
            return cls.load_docx(file_input)
        elif suffix in [".txt", ".md"]:
            return cls.load_txt(file_input)
        else:
            raise ValueError(f"Unsupported file format: '{suffix}' from filename '{filename}'")